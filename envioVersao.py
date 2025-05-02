import os
import zipfile
from docx import Document

def extract_images_from_docx(docx_file):
    # A extensão .docx é basicamente um arquivo ZIP
    with zipfile.ZipFile(docx_file, 'r') as docx_zip:
        # Filtrando os arquivos de imagem dentro do .docx
        image_files = [f for f in docx_zip.namelist() if f.startswith('word/media/')]
        
        # Criando o diretório para salvar as imagens extraídas
        if not os.path.exists("output_img"):
            os.makedirs("output_img")
        
        # Extraindo as imagens
        for image_file in image_files:
            image_data = docx_zip.read(image_file)
            image_name = os.path.join("output_img", os.path.basename(image_file))
            
            with open(image_name, 'wb') as img:
                img.write(image_data)

def lerdoc(template):
    extract_images_from_docx(docx_file=template)
    doc = Document(template)
    html = "<html><body>"

    html += f"<p>{doc.paragraphs[0].text}</p>"
    html += f"<p>{doc.paragraphs[1].text}</p>"
    html += f"<p>{doc.paragraphs[2].text}</p>"
    html += f"<p>{doc.paragraphs[3].text}</p>"
    pasta = os.path.join("output_img","\image1.png")

    html += f'<img src="{pasta}">'


    html += "</html></body>"
    return html



